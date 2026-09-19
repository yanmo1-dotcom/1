# -*- coding: utf-8 -*-
"""将游戏开发工程师简历内容生成为可编辑 Word 版（用于补齐个人信息后投递）"""
from pathlib import Path

from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[2]
OUT_DIR = ROOT / "output" / "pdf"
OUT = OUT_DIR / "游戏开发工程师简历.docx"
OUT_DIR.mkdir(parents=True, exist_ok=True)

NAVY = RGBColor(0x16, 0x28, 0x3B)
BLUE = RGBColor(0x1F, 0x6F, 0xB2)
MUTED = RGBColor(0x52, 0x64, 0x74)

doc = Document()

# 全局中文字体
style = doc.styles["Normal"]
style.font.name = "SimHei"
style.font.size = Pt(10.5)
style.element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")

def set_font(run, size=10.5, bold=False, color=NAVY, name="SimHei"):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color

def para(text="", size=10.5, bold=False, color=NAVY, align=None, space_after=4):
    p = doc.add_paragraph()
    if align:
        p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    r = p.add_run(text)
    set_font(r, size=size, bold=bold, color=color)
    return p

def bullet(text, size=10.5):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.left_indent = Cm(0.5)
    r = p.add_run("· " + text)
    set_font(r, size=size)
    return p

def section(title):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(title)
    set_font(r, size=14, bold=True, color=BLUE)
    # 下划线
    pPr = p._p.get_or_add_pPr()
    from docx.oxml import OxmlElement
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:color"), "D8E0E7")
    pBdr.append(bottom)
    pPr.append(pBdr)
    return p

# 标题
p = para("常天昊｜游戏开发工程师", size=20, bold=True, color=NAVY, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
para("求职方向：Java 游戏客户端 / 游戏逻辑开发（可胜任游戏服务端）", size=11, color=BLUE, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
para("电话：18516898857　｜　邮箱：cdl24523452@163.com　｜　城市：北京　｜　GitHub：github.com/yanmo1-dotcom/1",
     size=9.5, color=MUTED, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=8)

section("个人优势")
for t in [
    "独立开发 8 款可运行的桌面/网页游戏，覆盖沙盒、射击、卡牌、对战等品类，累计 200+ Java 源文件、近 3 万行代码，多数项目可在线试玩。",
    "熟悉游戏开发完整链路：游戏循环、渲染、物理与碰撞、状态机、敌人/Boss AI、对象池、粒子与音效、程序化地图、存档，及多人联机。",
    "掌握 LibGDX、LWJGL/OpenGL 图形接口与自研 Swing 渲染管线，具备渲染性能优化实践（滚动缓存、异步地形网格、BFS 光照）。",
    "独立实现局域网多人联机与 Web 实时对战（UDP 房间发现、TCP 主机权威同步、WebSocket/STOMP），可同时承担游戏客户端与服务端工作。",
    "关注玩法手感与性能细节（重力、惯性、难度曲线、FPS 稳定性），能从创意、编码、测试到网页演示独立交付完整可玩作品。",
]:
    bullet(t)

section("专业技能")
for t in [
    "语言基础：Java 17/21/26、OOP、集合、I/O、并发与多线程、Socket、算法与数据结构",
    "游戏引擎/图形：LibGDX、LWJGL 3、OpenGL、Swing 自绘渲染、GWT 网页版；滚动缓存渲染、异步网格生成、BFS 光照传播",
    "游戏系统：游戏循环、碰撞检测、物理积分、状态机、敌人/Boss 行为 AI、程序化地图与生物群系、对象池、粒子系统、音频、摄像机、UI/HUD、存档",
    "网络联机：TCP/UDP、主机权威模型、状态同步与插值、局域网房间发现、WebSocket/STOMP 实时对战",
    "后端与工具：Spring Boot、JDBC、H2/MySQL、JSON/CSV、Maven、Gradle 多模块、JUnit 5、Git",
    "游戏数学：程序化生成、噪声地形、空间检索、物理运动学、博弈与 DP 算法（LeetCode 实战）",
]:
    bullet(t)

section("重点项目经历")
def project(name, stack, items):
    para(name, size=12, bold=True, space_after=1)
    para("技术栈：" + stack, size=9.5, color=BLUE, space_after=2)
    for it in items:
        bullet(it, size=10.5)

project("泰拉瑞亚 Lite｜横版沙盒游戏 + 局域网联机（个人项目）",
        "Java（JDK 26）、Swing、TCP/UDP、多线程，无第三方依赖",
        [
            "实现“挖掘→收集→建造→合成→战斗”完整玩法闭环；程序化生成森林/沙漠/雪原/丛林四大生物群系与平原、丘陵、山地地形。",
            "实现史莱姆/僵尸行为 AI 与克苏鲁之眼 Boss（飞行追踪、周期冲刺、半血狂暴、顶部血条），完成近战挥砍弧光与远程弓重力弹道。",
            "设计主机权威联机模型：UDP 广播房间发现、TCP 100ms 状态同步与远端平滑插值、增量方块事件、全服聊天，支持 5 人同场。",
            "实现方块层滚动缓存渲染与可见范围光照，实测移动/持续战斗下 FPS 稳定 170~210，显著降低卡顿。",
            "编写 130+ 项断言的无 GUI 自测（逻辑、联机协议、地形生成、Boss、背包、局域网发现）并全部通过。",
        ])

project("体素泰拉瑞亚｜LWJGL/OpenGL 3D 体素生存（个人项目）",
        "Java 17、LWJGL 3、OpenGL、Gradle",
        [
            "实现种子地形生成、区块流式加载/卸载、方块放置与破坏、碰撞与射线拾取；固定线程池异步生成地形与网格、主线程限帧上传。",
            "实现六大 Boss 多阶段 AI、NPC 入住判定、血月/雨天/哥布林入侵等世界事件，以及木石铁金→首领装→暗影/寒霜/蜂甲的完整装备进阶链。",
            "局域网房主权威：客机挖掘/放置经请求-校验-写入，怪物 AI 与伤害由房主裁定，战利品归属同步给攻击者。",
            "实现太阳光与点光源 BFS 传播、跨区块访问，并在区块卸载时正确释放 OpenGL 资源。",
        ])

project("Kapai Roguelike 卡牌战斗游戏｜libGDX 多模块（个人项目）",
        "Java 17、LibGDX、Maven 多模块、Jackson、JUnit 5",
        [
            "拆分 core/data/ui 三模块，隔离战斗逻辑、数据加载与界面渲染，便于独立测试与替换。",
            "以状态机组织玩家回合、怪物意图、行动与结算，实现牌堆、能量、护盾、状态效果与目标选择。",
            "使用效果工厂组合伤害/格挡/抽牌/状态附加，通过监听器驱动遗物触发；JSON 驱动卡牌配置，集成测试覆盖数据加载、出牌伤害与遗物效果。",
        ])

project("奇斯要塞｜libGDX 俯视角策略射击（个人项目）",
        "Java、LibGDX、Gradle、GWT",
        [
            "实现步兵/吉普/坦克多载具、敌人与 Boss、武器状态机、TileMap 关卡与计分系统。",
            "使用对象池管理弹体与特效，配以粒子与音效反馈；第二版通过 GWT 编译为网页可玩版本，浏览器即点即玩。",
        ])

project("飞机射击 Racing Game｜Swing 弹幕射击（个人项目）",
        "Java、Swing、音频",
        ["实现多类型敌机、Boss、波次管理、道具掉落、粒子爆炸与音效反馈，含本地最高分存档。"])

project("冰球游戏｜Swing 局域网对战（个人项目）",
        "Java、Swing、Socket、Maven",
        [
            "完成单人 AI 与局域网双人模式，含菜单、计分、胜负、重开与音效；房主统一计算球运动与比分保证画面一致。",
            "网络读写置于守护线程，并通过 Swing 事件线程安全更新界面。",
        ])

section("其他项目")
for t in [
    "Web 实时聊天与在线对战平台：Spring Boot + WebSocket/STOMP，接入五子棋/井字棋/猜数字/石头剪刀布，服务端校验动作并裁决胜负。",
    "A 股实时行情网站：Spring Boot + OkHttp 拉取行情接口，定时刷新内存数据并提供多维排行。",
    "命令行多房间聊天室：多线程 Socket 实现服务端、客户端、多房间管理与日志服务。",
    "LeetCode 算法练习：Java 完成石子游戏、最小回文排列、图与字符串等题目（博弈、DP、贪心）。",
]:
    bullet(t)

section("教育经历")
para("河北司法警官职业学院｜罪犯心理测量与矫正技术｜大专｜2023.09 - 2026.06")

section("自我评价")
para("对游戏开发有持续热情，能从想法、玩法设计、编码、测试到网页演示独立完成一款可玩的小游戏；注重手感与性能细节，善于通过可运行项目快速学习新引擎与新技术。希望进入游戏行业，参与从原型到上线的完整流程，并快速补足工程规范与团队协作经验。")

doc.save(str(OUT))
print(OUT)
